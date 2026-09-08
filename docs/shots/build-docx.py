# -*- coding: utf-8 -*-
"""
화면 안내 문서를 워드(.docx)로 만든다. 글만 있는 것과 사진이 들어간 것 둘.

    python docs/shots/build-docx.py

`보고서-화면안내.md` 하나를 읽어 둘 다 만든다. 글은 완전히 같고 사진 유무만
다르다. 문구를 고칠 때는 `.md` 만 고치고 이걸 다시 돌리면 된다.

⚠️ 워드로 파일을 열어 둔 채로 돌리면 「Permission denied」 로 막힌다.
   닫고 다시 돌려야 한다.

사진을 어느 제목 뒤에 넣을지는 `build.mjs` 와 같은 목록을 쓴다. 두 곳에
따로 적으면 한쪽만 고쳤을 때 어긋난다.
"""

import io
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
DOCS = os.path.dirname(HERE)
SRC = os.path.join(DOCS, "보고서-화면안내.md")

# build.mjs 의 SHOTS 에서 그대로 읽어 온다 — 두 곳에 적지 않는다
def load_shots():
    js = io.open(os.path.join(HERE, "build.mjs"), encoding="utf-8").read()
    block = js[js.index("const SHOTS = ["): js.index("];", js.index("const SHOTS = ["))]
    out = []
    for m in re.finditer(r'\["([^"]*)",\s*"([^"]+)",\s*"([^"]+)"\]', block):
        out.append((m.group(1), m.group(2), m.group(3)))
    return out

SHOTS = load_shots()

# 제목 → 그 뒤에 넣을 (파일, 설명) 목록
BY_HEADING = {}
_current = None
for heading, file, caption in SHOTS:
    if heading:
        _current = heading
        BY_HEADING[_current] = []
    if _current:
        BY_HEADING[_current].append((file, caption))


INK = RGBColor(0x1A, 0x19, 0x17)
MUTED = RGBColor(0x6B, 0x66, 0x60)


def style(doc):
    """본문 글꼴과 줄 간격. 워드 기본값은 한글에 좁다."""
    n = doc.styles["Normal"]
    n.font.name = "맑은 고딕"
    n.font.size = Pt(10.5)
    n.font.color.rgb = INK
    # ⚠️ 한 문단 **안**의 줄 간격은 좁게, 문단 **사이**는 넓게 (2026-08-26).
    # 1.5로 두었더니 한 문단이 뜯어져 보였다. 줄이 멀면 이어지는 문장인지
    # 새 이야기인지 눈이 못 가른다. 문단 구분은 사이 여백이 맡는다.
    n.paragraph_format.space_after = Pt(11)
    n.paragraph_format.line_spacing = 1.25
    # 한글 글꼴은 동아시아 속성에도 따로 넣어야 적용된다
    rpr = n.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    from docx.oxml.ns import qn
    rfonts.set(qn("w:eastAsia"), "맑은 고딕")

    # 제목 크기 차이를 크게 두지 않는다. 꾸미는 것보다 어디가 끊기는지가 중요하다
    for name, size, before in [
        ("Heading 1", 15, 22),
        ("Heading 2", 12.5, 18),
        ("Heading 3", 11, 14),
    ]:
        s = doc.styles[name]
        s.font.name = "맑은 고딕"
        s.font.size = Pt(size)
        s.font.color.rgb = INK
        s.font.bold = True
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(6)
        rpr = s.element.get_or_add_rPr()
        rpr.get_or_add_rFonts().set(qn("w:eastAsia"), "맑은 고딕")


def add_rich(p, text):
    """**굵게** 와 `코드` 만 처리한다. 그 밖은 그대로 둔다."""
    for part in re.split(r"(\*\*.+?\*\*|`.+?`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        else:
            p.add_run(part)


def add_images(doc, heading):
    for file, caption in BY_HEADING.get(heading, []):
        path = os.path.join(HERE, file)
        if not os.path.exists(path):
            print("  사진 없음:", file)
            continue
        doc.add_picture(path, width=Cm(15))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(caption)
        r.font.size = Pt(9)
        r.font.color.rgb = MUTED


def build(with_images, out_name):
    doc = Document()
    style(doc)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.2)
        s.left_margin = s.right_margin = Cm(2.5)

    md = io.open(SRC, encoding="utf-8").read()

    # 접속 주소는 커밋되지 않는 site-url.txt 에 있다 (build.mjs 와 같은 이유).
    # 파일이 없으면 주소 없이 만들어진다.
    url_file = os.path.join(HERE, "site-url.txt")
    if os.path.exists(url_file):
        url = io.open(url_file, encoding="utf-8").read().strip()
        nl = chr(10)
        at = md.find(nl)  # 첫 제목 줄 끝
        if url and at > 0:
            md = md[: at + 1] + nl + url + nl + md[at + 1 :]
    for raw in md.split("\n"):
        line = raw.rstrip()

        if line == "---":
            continue  # 워드에서는 가로줄 대신 제목 앞 여백으로 나눈다

        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
            if with_images:
                add_images(doc, line)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
            if with_images:
                add_images(doc, line)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_rich(p, line[2:])
        elif line == "":
            continue
        else:
            p = doc.add_paragraph()
            add_rich(p, line)

    out = os.path.join(DOCS, out_name)
    doc.save(out)
    size = os.path.getsize(out) / 1024
    print("  만들었습니다: %s  (%.0fKB)" % (out_name, size))


build(False, "화면안내-글만.docx")
build(True, "화면안내-사진포함.docx")
